"""The AI vs Human report as an Excel workbook (.xlsx), one sheet per part, for filtering and pivoting.

Sheets: Verdict · Improvement plan · Root causes · Root cause calls · Call RCA · Recommended changes ·
Scores per call (same columns as the CSV).
"""

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="F1F2F4")


def _sheet(wb: Workbook, title: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title)
    if not rows:
        ws.append(["Nothing here for this selection."])
        return
    headers = list(rows[0])
    ws.append(headers)
    for r in rows:
        ws.append([_cell(r.get(h)) for h in headers])
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill = Font(bold=True), HEAD_FILL
        width = max(
            [len(str(h))]
            + [len(str(ws.cell(row=n, column=i).value or "")) for n in range(2, min(ws.max_row, 200) + 1)]
        )
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _cell(v: Any) -> Any:
    if isinstance(v, list):
        return "; ".join(str(x) for x in v)
    if isinstance(v, dict):
        return str(v)
    return v


def render_xlsx(result: dict[str, Any], scores_per_call: list[dict[str, Any]]) -> bytes:
    imp, syn = result.get("improvement") or {}, result.get("synthesis") or {}
    v = imp.get("verdict") or {}
    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]

    verdict = [
        {"item": "Verdict", "value": v.get("label")},
        {"item": "Bot reaches % of human quality", "value": v.get("readiness_pct")},
        {"item": "Bot average review score (1-5)", "value": result["scores"]["ai"]["avg_review"]},
        {"item": "Human average review score (1-5)", "value": result["scores"]["human"]["avg_review"]},
        {"item": "Bot calls ending with a concrete next step (%)", "value": (v.get("positive_outcome_pct") or {}).get("ai")},
        {"item": "Human calls ending with a concrete next step (%)", "value": (v.get("positive_outcome_pct") or {}).get("human")},
        {"item": "Fix first", "value": v.get("top_levers")},
        {"item": "Headline (AI)", "value": syn.get("verdict_headline") or result.get("synthesis_error")},
        {"item": "Assessment (AI)", "value": syn.get("verdict_detail")},
        {"item": "Report built", "value": result.get("generated_at")},
    ] + [{"item": f"Key point {i + 1}", "value": p} for i, p in enumerate(v.get("points", []))]  # fmt: skip
    _sheet(wb, "Verdict", verdict)
    _sheet(wb, "Improvement plan", [
        {"rank": i + 1, "priority": p["priority"], "parameter": p["label"], "bot": p["ai"], "humans": p["human"],
         "gap": p["gap"], "weak bot calls": p["weak_calls"], "of bot calls": p["of"], "owner": p["area"],
         "how to fix": p["fix"]}
        for i, p in enumerate(imp.get("parameters", []))
    ])  # fmt: skip
    _sheet(wb, "Root causes", [
        {"failure": c["label"], "bot calls": c["count"], "of": c["of"], "share %": c["share"],
         "priority": c["priority"], "owner": c["area"], "fix": c["fix"]}
        for c in imp.get("root_causes", [])
    ])  # fmt: skip
    _sheet(wb, "Root cause calls", [
        {"failure": c["label"], "call id": h["call_id"], "call": h["label"], "lead id": h["lead_id"],
         "what happened": h["description"], "quote": h["quote"], "at second": h["t"], "better": h["better"]}
        for c in imp.get("root_causes", []) for h in c["calls"]
    ])  # fmt: skip
    _sheet(wb, "Call RCA", [
        {"call id": x["call_id"], "call": x["label"], "lead id": x["lead_id"], "call score": x["review_score"],
         "outcome": x["outcome"], "at second": i["t"], "kind": i["kind"], "issue": i["title"],
         "what happened": i["detail"], "quote": i["quote"], "better": i["better"]}
        for x in imp.get("call_rca", []) for i in x["issues"]
    ])  # fmt: skip
    _sheet(wb, "Recommended changes", [
        {"priority": ch["priority"], "area": ch.get("area"), "change": ch["change"], "why": ch["rationale"],
         "new bot line": ch.get("bot_line"), "example calls": [e["label"] for e in ch.get("examples", [])]}
        for ch in sorted(syn.get("recommended_changes", []), key=lambda c: ["high", "medium", "low"].index(c["priority"]))
    ])  # fmt: skip
    _sheet(wb, "Scores per call", scores_per_call)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
