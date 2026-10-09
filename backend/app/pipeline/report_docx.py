"""Render the report outline (report.py) as a Word document (.docx) with python-docx.

Word handles Hindi (Devanagari) quotes itself, so no special fonts are needed here.
"""

import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.pipeline.report_blocks import Bullets, Para, Quote, Report, Sub, Table

MUTED = RGBColor(0x6B, 0x72, 0x80)
AI_BLUE = RGBColor(0x2F, 0x6F, 0xDE)
HUMAN_ORANGE = RGBColor(0xE8, 0x69, 0x2C)


def _shade(cell, hex_fill: str) -> None:  # type: ignore[no-untyped-def]
    """Background colour for a table cell (python-docx has no direct API for this)."""
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    props.append(shd)


def _table(doc, block: Table) -> None:  # type: ignore[no-untyped-def]
    t = doc.add_table(rows=1, cols=len(block.headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(block.headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        if h == "AI voice bot":
            run.font.color.rgb = AI_BLUE
        elif h == "Human agents":
            run.font.color.rgb = HUMAN_ORANGE
        _shade(cell, "F1F2F4")
    for row in block.rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    doc.add_paragraph()


def render_docx(report: Report) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    doc.add_heading(report.title, level=0)
    for line in report.meta:
        p = doc.add_paragraph()
        r = p.add_run(line)
        r.font.color.rgb = MUTED

    for section in report.sections:
        doc.add_heading(section.heading, level=1)
        for block in section.blocks:
            if isinstance(block, Sub):
                doc.add_heading(block.text, level=2)
            elif isinstance(block, Para):
                run = doc.add_paragraph().add_run(block.text)
                if block.muted:
                    run.font.color.rgb = MUTED
            elif isinstance(block, Bullets):
                for item in block.items:
                    doc.add_paragraph(item, style="List Bullet")
            elif isinstance(block, Quote):
                p = doc.add_paragraph(style="Intense Quote")
                p.add_run(f"“{block.text}”")
                src = doc.add_paragraph().add_run(block.source)
                src.font.size = Pt(9)
                src.font.color.rgb = MUTED
            elif isinstance(block, Table):
                _table(doc, block)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
