"""The shareable AI vs Human report (Word + print page) built from one outline."""

import io

from docx import Document

from app.pipeline.report import Quote, build_outline
from app.pipeline.report_docx import render_docx
from app.pipeline.report_html import render_html

SIDE = {"uploaded": 2, "analysed": 2, "failed": 0, "not_connected": 0, "in_progress": 0, "audio_seconds": 600}
SIGNALS = {"mood_end": {"positive": 1}, "mood_improved": 1, "mood_worsened": 0, "objections_total": 1,
           "objections_handled_well": 1, "friction_per_call": 0.5, "unanswered_per_call": 0, "escalated_pct": 0,
           "avg_intent": 60}  # fmt: skip
RESULT = {
    "generated_at": "2026-10-09T06:30:00+00:00",
    "scope": {"batch_ids": [3], "campaign": "", "date_from": None, "date_to": None, "comparable_only": False},
    "records": {"ai": SIDE, "human": SIDE},
    "sample_warning": "Based on 2 AI and 2 human calls.",
    "scores": {"ai": {"avg_review": 3.0, "n": 2, "of": 2}, "human": {"avg_review": 3.3, "n": 2, "of": 2}},
    "dimensions": [{"key": "clarity", "label": "Clarity", "ai": 3.0, "human": 3.5, "ai_n": 2, "human_n": 2}],
    "metrics": [{"key": "duration_s", "label": "Call length (s)", "ai": 95.0, "human": 130.0}],
    "outcomes": [{"key": "store_visit", "label": "Store visit agreed", "ai": 0, "human": 50}],
    "signals": {"ai": SIGNALS, "human": SIGNALS},
    "improvement": {
        "verdict": {"status": "behind", "label": "The bot is behind human agents", "gap": 0.3, "readiness_pct": 91,
                    "positive_outcome_pct": {"ai": 50, "human": 100}, "top_levers": ["Clarity"],
                    "points": ["The bot averages 3.0 / 5 against 3.3 / 5 for human agents (91% of human quality)."]},
        "parameters": [{"key": "clarity", "label": "Clarity", "area": "Script", "fix": "Shorter turns", "ai": 3.0,
                        "human": 3.5, "target": 3.5, "gap": 0.5, "weak_calls": 1, "of": 2, "weak_share": 50,
                        "priority": "high", "status": "behind",
                        "bot_examples": [{"call_id": 1, "label": "A1", "lead_id": 1, "score": 2, "reason": "Long",
                                          "quote": "Hello! I am calling", "t": 3, "better": "Say the price first"}],
                        "human_examples": []}],
        "root_causes": [{"pattern": "scripted_repeat", "label": "Repeats the same scripted line", "area": "Conversation flow",
                         "fix": "Rephrase", "count": 1, "of": 2, "share": 50, "priority": "high",
                         "calls": [{"call_id": 1, "label": "A1", "lead_id": 1, "description": "Repeated intro",
                                    "quote": "Hello! I am calling", "t": 3, "better": "Ji, batayiye"}]}],
        "objections": [],
        "call_rca": [{"call_id": 1, "label": "A1", "lead_id": 1, "duration_s": 95, "review_score": 2.6,
                      "outcome": "callback_scheduled", "mood_end": "neutral", "one_liner": "Price not given",
                      "main_issue": "Repeats the same scripted line", "issue_count": 1,
                      "issues": [{"kind": "failure", "title": "Repeats the same scripted line", "detail": "Repeated intro",
                                  "quote": "Hello! I am calling", "t": 3, "better": "Ji, batayiye"}]}],
    },
    "calls": {"ai": [{"id": 1, "label": "A1", "duration_s": 95, "status": "done", "error_code": None, "language": "hinglish",
                      "outcome": "callback_scheduled", "quality_pct": 40.0, "review_score": 2.6, "one_liner": "Price <not> given"}],
              "human": []},
    "synthesis": {
        "verdict_headline": "Humans secure more concrete next steps.",
        "differences": [{"theme": "Answering price", "ai": "Repeats script", "human": "Gives price", "why": "Price first",
                         "evidence": [{"call_id": 2, "quote": "के ये भाई साढ़े छह का मैट्रिस चाहिए", "t": 50,
                                       "label": "+91-9437612631", "agent_type": "ai"}]}],
        "ai_better": ["Consistent greeting"], "human_better": ["Adapts"], "outcomes_paragraph": "Humans book visits.",
        "recommended_changes": [{"priority": "low", "area": "Script", "change": "Shorter intro", "rationale": "Time",
                                 "bot_line": "", "examples": []},
                                {"priority": "high", "area": "Knowledge base", "change": "Answer price first",
                                 "rationale": "Most asked", "bot_line": "Ji, offer price 24,999 hai",
                                 "examples": [{"call_id": 1, "label": "A1", "lead_id": 1}]}],
    },
    "synthesis_error": None,
}  # fmt: skip


def test_outline_has_every_section_in_page_order():
    outline = build_outline(RESULT)
    assert [s.heading for s in outline.sections] == [
        "Records", "Verdict", "Improvement plan for the bot", "Root causes — repeated bot failures",
        "Recommended changes to the bot", "Call-by-call RCA (bot calls)", "Where they differ",
        "Review scores (1 = poor … 5 = excellent)", "Measured from the transcripts", "Outcomes and customer mood",
        "Every call",
    ]  # fmt: skip
    quotes = {b.source: b for s in outline.sections for b in s.blocks if isinstance(b, Quote)}
    assert quotes["AI voice bot · +91-9437612631 · 0:50"].href == "/calls/2?t=50"  # drill-down link
    assert "AI voice bot · A1 · 0:03 · scored 2/5 — Long" in quotes


def test_word_document_keeps_hindi_and_orders_changes_by_priority():
    doc = Document(io.BytesIO(render_docx(build_outline(RESULT))))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "के ये भाई साढ़े छह का मैट्रिस चाहिए" in text  # Devanagari survives
    change_rows = [r.cells[0].text for t in doc.tables for r in t.rows if r.cells[0].text in ("High", "Low")]
    assert change_rows == ["High", "Low"]  # most important first
    cells = {c.text for t in doc.tables for r in t.rows for c in r.cells}
    assert "“Ji, offer price 24,999 hai”" in cells and "A1" in cells  # new bot line + example call
    assert "Better: Say the price first" in text


def test_print_page_escapes_text_and_can_open_print_dialog():
    page = render_html(build_outline(RESULT), auto_print=True)
    assert "Price &lt;not&gt; given" in page  # escaped, never raw HTML
    assert '<a href="/calls/1?t=3">' in page  # quotes open the call at that moment
    assert "window.print()" in page and "@page" in page
    assert "<script>" not in render_html(build_outline(RESULT), auto_print=False).split("</style>")[1]


def test_pdf_is_made_on_the_server_with_hindi_text():
    from app.pipeline.report_pdf import render_pdf

    data = render_pdf(build_outline(RESULT), "https://app.example.com")
    assert data.startswith(b"%PDF") and len(data) > 5000
    assert b"https://app.example.com/calls/1?t=3" in data  # quote sources link to the call moment


def test_excel_has_one_sheet_per_part():
    from openpyxl import load_workbook

    from app.pipeline.report_xlsx import render_xlsx

    wb = load_workbook(io.BytesIO(render_xlsx(RESULT, [{"call_id": 1, "label": "A1"}])))
    assert wb.sheetnames == ["Verdict", "Improvement plan", "Root causes", "Root cause calls", "Call RCA",
                             "Recommended changes", "Scores per call"]  # fmt: skip
    assert wb["Recommended changes"]["E2"].value == "Ji, offer price 24,999 hai"
